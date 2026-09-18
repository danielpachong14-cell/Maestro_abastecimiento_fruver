"""Auditoría: combinaciones (tienda × ítem) SIN DATOS en Celes.

Problema que audita
-------------------
En el Output diario aparecen filas con Consumo Diario = 0, Stock Antes Pedido = 0
y Días vacíos, que reciben Pedido Final = 1. Esas combinaciones tienda-ítem NO
tuvieron match en Celes: el preprocessor las rellena con consumo=0/inventario=0,
lo que las clasifica como AGOTADA sin estarlo y les despacha una caja mínima. El
riesgo de negocio es surtir a tiendas que quizá no lo necesitan solo porque falta
el dato en Celes.

Qué entrega
-----------
1. VISTA HOY (exacta): reproduce el pipeline real (load_files +
   build_distribution_df) y lista cada par con `sin_match_celes == True`,
   diagnosticando la causa (tienda ausente de Celes / ítem ausente / par
   faltante). Es la fuente de verdad, no una inferencia.
2. VISTA HISTÓRICA (proxy): recorre los Output/distribucion_fruver_*.xlsx ya
   generados y cuenta la frecuencia de la "firma sin datos" por (tienda × ítem),
   por ítem y por tienda. Es aproximada (solo hay un Celes.xlsx, el de hoy; no
   hay históricos que reproducir con fidelidad).

Salida: Output/AUDITORIA_SinDatos_Celes.xlsx (no modifica ningún otro archivo).

Uso:
    cd Distribucion-Fruver
    distribuidor-fruver/venv/Scripts/python.exe auditoria_sin_datos_celes.py
"""

import glob
import os
import re
import sys
from collections import defaultdict
from contextlib import ExitStack
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).parent
INPUT = ROOT / "Input"
OUTPUT = ROOT / "Output"

sys.path.insert(0, str(ROOT / "distribuidor-fruver"))

from core.loader import load_files                       # noqa: E402
from core.preprocessor import (                          # noqa: E402
    build_distribution_df,
    normalize_store,
    _normalize_item_series,
)

# Mismo mapeo de rutas que run.py.
FILES = {
    "stock":        INPUT / "Stock.xlsx",
    "celes":        INPUT / "Celes.xlsx",
    "portafolio":   INPUT / "DB_Portafolio Fruver.xlsx",
    "tiendas":      INPUT / "DB_Tiendas.xlsx",
    "tiendas_item": INPUT / "DB_Tiendas Por itmes y portafolio.xlsx",
    "espejo":       INPUT / "DB_ProductosEspejo.xlsx",
}
EXCLUIDAS_PATH = INPUT / "Tiendas_No generar pedido.xlsx"

# Índices de columna de la hoja 'Distribución' (header de 12 columnas), verificados
# contra los Outputs reales:
#   0 Centro Operacional (store)   1 Nombre Tienda   2 Zona   3 Código de Producto
#   4 Nombre Ítem   5 UM   6 Consumo Diario   7 Stock Antes Pedido
#   8 Días Inventario Actual   9 Pedido Final   10 Stock Después   11 Días Proyectado
C_STORE, C_TIENDA, C_ZONA, C_ITEM, C_INAME = 0, 1, 2, 3, 4
C_CONSUMO, C_STOCK_ANTES, C_DIAS_ACT, C_PEDIDO = 6, 7, 8, 9

SHEET_DISTRIB = "Distribución"  # 'Distribución'
AUDIT_PATH = OUTPUT / "AUDITORIA_SinDatos_Celes.xlsx"


def _num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------
# VISTA HOY (exacta) — reproduce el pipeline
# ---------------------------------------------------------------------------
def vista_hoy():
    """Corre load_files + build_distribution_df sobre el Input actual y devuelve
    el DataFrame de pares sin match con la causa diagnosticada."""
    with ExitStack() as stack:
        archivos = [stack.enter_context(open(p, "rb")) for p in FILES.values()]
        excl = (
            stack.enter_context(open(EXCLUIDAS_PATH, "rb"))
            if EXCLUIDAS_PATH.exists() else None
        )
        dfs = load_files(*archivos, tiendas_excluidas=excl)

    df_merged, _ = build_distribution_df(dfs)
    sm = df_merged[df_merged["sin_match_celes"]].copy()

    # Diagnóstico de causa: normalizar el Celes crudo con las MISMAS funciones que
    # usa el preprocessor, para ver qué tiendas/ítems/pares sí existen en Celes.
    celes = dfs["celes"].copy()
    celes["cod_siesa"] = celes["Código de Bodega"].apply(normalize_store)
    celes["item_key"] = _normalize_item_series(celes["Código de Producto"])
    celes = celes.dropna(subset=["cod_siesa", "item_key"])
    celes["item_key"] = celes["item_key"].astype(int)

    tiendas_en_celes = set(celes["cod_siesa"])
    items_en_celes = set(celes["item_key"])

    def causa(row):
        st, it = row["store_code"], int(row["item_code"])
        if st not in tiendas_en_celes:
            return "Tienda ausente de Celes"
        if it not in items_en_celes:
            return "Ítem ausente de todo Celes"
        return "Par tienda-ítem faltante"

    if len(sm):
        sm["causa"] = sm.apply(causa, axis=1)
    else:
        sm["causa"] = pd.Series(dtype=str)

    cols = [
        "store_code", "store_name", "zona", "item_code", "item_desc", "um",
        "cajas_disponibles_cedi", "grupo_id", "causa",
    ]
    return sm[cols].sort_values(["item_code", "store_code"]).reset_index(drop=True), len(df_merged)


# ---------------------------------------------------------------------------
# VISTA HISTÓRICA (proxy) — recorre los Outputs
# ---------------------------------------------------------------------------
def vista_historica():
    files = sorted(glob.glob(str(OUTPUT / "distribucion_fruver_*.xlsx")))
    files = [f for f in files if "~$" not in os.path.basename(f)]

    sindatos = defaultdict(list)   # (store, item) -> [fechas]
    present = defaultdict(set)     # (store, item) -> {fechas presentes en Distribución}
    item_name, store_name, zona = {}, {}, {}
    conteo_oficial = {}            # fecha -> combos sin match (hoja Alertas)

    for f in files:
        m = re.search(r"(\d{8})", os.path.basename(f))
        if not m:
            continue
        fecha = m.group(1)
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws = wb[SHEET_DISTRIB] if SHEET_DISTRIB in wb.sheetnames else wb.worksheets[0]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[C_STORE] is None or r[C_ITEM] is None:
                continue
            store, item = r[C_STORE], r[C_ITEM]
            key = (store, item)
            present[key].add(fecha)
            item_name[item] = r[C_INAME]
            store_name[store] = r[C_TIENDA]
            zona[store] = r[C_ZONA]
            consumo, stock_antes = _num(r[C_CONSUMO]), _num(r[C_STOCK_ANTES])
            dias = r[C_DIAS_ACT]
            if consumo == 0 and stock_antes == 0 and (dias is None or dias == ""):
                sindatos[key].append(fecha)

        # Conteo oficial desde la hoja Alertas (columna 'Combos Sin Match').
        if "Alertas" in wb.sheetnames:
            wa = wb["Alertas"]
            hdr = next(wa.iter_rows(max_row=1, values_only=True), ())
            hdr = [str(h) for h in hdr]
            idx = hdr.index("Combos Sin Match") if "Combos Sin Match" in hdr else None
            for row in wa.iter_rows(min_row=2, values_only=True):
                texto = " ".join(str(c) for c in row if c is not None)
                if "match" in texto.lower() and "celes" in texto.lower():
                    val = row[idx] if idx is not None and idx < len(row) else None
                    if isinstance(val, (int, float)):
                        conteo_oficial[fecha] = int(val)
        wb.close()

    return {
        "files": files,
        "sindatos": sindatos,
        "present": present,
        "item_name": item_name,
        "store_name": store_name,
        "zona": zona,
        "conteo_oficial": conteo_oficial,
    }


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------
def _autofit(ws, max_width=60):
    for col in ws.columns:
        length = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = min(length + 2, max_width)


def _bold_header(ws):
    for c in ws[1]:
        c.font = openpyxl.styles.Font(bold=True)


def exportar(hoy_df, total_hoy, hist):
    wb = openpyxl.Workbook()

    # --- Hoja 1: Hoy - Detalle ---
    ws = wb.active
    ws.title = "Hoy - Detalle"
    ws.append([
        "Código Tienda", "Nombre Tienda", "Zona", "Código Ítem",
        "Nombre Ítem", "UM", "Cajas CEDI", "Grupo Espejo", "Causa",
    ])
    for _, r in hoy_df.iterrows():
        ws.append([
            r["store_code"], r["store_name"], r["zona"], int(r["item_code"]),
            r["item_desc"], r["um"], int(r["cajas_disponibles_cedi"]),
            r["grupo_id"], r["causa"],
        ])
    _bold_header(ws)
    _autofit(ws)

    # --- Hoja 2: Hoy - Resumen ---
    ws = wb.create_sheet("Hoy - Resumen")
    ws.append(["POR ÍTEM (hoy)"])
    ws.append(["Código Ítem", "Nombre Ítem", "# Tiendas sin datos"])
    if len(hoy_df):
        por_item = hoy_df.groupby(["item_code", "item_desc"]).size().sort_values(ascending=False)
        for (item, desc), n in por_item.items():
            ws.append([int(item), desc, int(n)])
    ws.append([])
    ws.append(["POR TIENDA (hoy)"])
    ws.append(["Código Tienda", "Nombre Tienda", "# Ítems sin datos"])
    if len(hoy_df):
        por_tnd = hoy_df.groupby(["store_code", "store_name"]).size().sort_values(ascending=False)
        for (st, nm), n in por_tnd.items():
            ws.append([st, nm, int(n)])
    _autofit(ws)

    sindatos = hist["sindatos"]
    present = hist["present"]
    item_name = hist["item_name"]
    store_name = hist["store_name"]
    zona = hist["zona"]

    # --- Hoja 3: Histórico - Pares ---
    ws = wb.create_sheet("Histórico - Pares")
    ws.append([
        "Código Tienda", "Nombre Tienda", "Zona", "Código Ítem",
        "Nombre Ítem", "# Fechas Sin Datos", "# Fechas Presente",
        "% Sin Datos", "Fechas",
    ])
    for (store, item), fechas in sorted(sindatos.items(), key=lambda kv: -len(kv[1])):
        npres = len(present[(store, item)])
        pct = round(100 * len(fechas) / npres, 1) if npres else 0.0
        ws.append([
            store, store_name.get(store, ""), zona.get(store, ""), item,
            item_name.get(item, ""), len(fechas), npres, pct,
            ", ".join(sorted(fechas)),
        ])
    _bold_header(ws)
    _autofit(ws)

    # --- Hoja 4: Histórico - Ítem ---
    ws = wb.create_sheet("Histórico - Ítem")
    ws.append([
        "Código Ítem", "Nombre Ítem", "# Tiendas", "# Fechas",
        "Total Ocurrencias",
    ])
    it_st, it_fx, it_oc = defaultdict(set), defaultdict(set), defaultdict(int)
    for (store, item), fechas in sindatos.items():
        it_st[item].add(store)
        it_fx[item].update(fechas)
        it_oc[item] += len(fechas)
    for item in sorted(it_oc, key=lambda i: -it_oc[i]):
        ws.append([item, item_name.get(item, ""), len(it_st[item]),
                   len(it_fx[item]), it_oc[item]])
    _bold_header(ws)
    _autofit(ws)

    # --- Hoja 5: Histórico - Tienda ---
    ws = wb.create_sheet("Histórico - Tienda")
    ws.append([
        "Código Tienda", "Nombre Tienda", "# Ítems", "# Fechas",
        "Total Ocurrencias",
    ])
    st_it, st_fx, st_oc = defaultdict(set), defaultdict(set), defaultdict(int)
    for (store, item), fechas in sindatos.items():
        st_it[store].add(item)
        st_fx[store].update(fechas)
        st_oc[store] += len(fechas)
    for store in sorted(st_oc, key=lambda s: -st_oc[s]):
        ws.append([store, store_name.get(store, ""), len(st_it[store]),
                   len(st_fx[store]), st_oc[store]])
    _bold_header(ws)
    _autofit(ws)

    # --- Hoja 6: Histórico - Fechas (proxy vs oficial) ---
    ws = wb.create_sheet("Histórico - Fechas")
    ws.append([
        "Fecha", "Proxy (firma sin datos)", "Oficial (hoja Alertas)",
        "Diferencia", "Nota",
    ])
    por_fecha_proxy = defaultdict(int)
    for (store, item), fechas in sindatos.items():
        for fch in fechas:
            por_fecha_proxy[fch] += 1
    oficial = hist["conteo_oficial"]
    todas = sorted(set(por_fecha_proxy) | set(oficial))
    for fch in todas:
        p = por_fecha_proxy.get(fch)
        o = oficial.get(fch)
        if o is None:
            nota = "Output sin instrumentación de sin-match"
            diff = ""
        elif p is None:
            nota = "Sin filas proxy en Distribución"
            diff = o
        else:
            diff = o - p
            if diff > 0:
                nota = "Proxy subestima (sin-match sin pedido no aparece en Distribución)"
            elif diff < 0:
                nota = "Proxy sobreestima (incluye tiendas con consumo genuino 0, no sin-match)"
            else:
                nota = "Proxy = Oficial"
        ws.append([fch, p if p is not None else "", o if o is not None else "", diff, nota])
    _bold_header(ws)
    _autofit(ws)

    # --- Hoja 7: Metodología ---
    ws = wb.create_sheet("Metodología")
    notas = [
        "AUDITORÍA — Combinaciones (tienda × ítem) SIN DATOS en Celes",
        "",
        "QUÉ SIGNIFICA 'SIN DATOS'",
        "  La combinación tienda-ítem es elegible (está en el portafolio activo y",
        "  la tienda está activa) pero NO tiene fila en Celes. El preprocessor la",
        "  rellena con consumo=0 e inventario=0 -> se clasifica AGOTADA y recibe una",
        "  caja mínima, aunque quizá no esté realmente agotada.",
        "",
        "VISTA HOY (hojas 'Hoy - ...') — EXACTA",
        "  Se reproduce el pipeline real (core.loader.load_files +",
        "  core.preprocessor.build_distribution_df) y se lee la columna",
        "  'sin_match_celes' que el propio preprocessor calcula. Es la fuente de",
        "  verdad, NO una inferencia. La columna 'Causa' distingue:",
        "    - Tienda ausente de Celes: Celes no reporta esa bodega para nada.",
        "    - Ítem ausente de todo Celes: ese código no aparece en ninguna fila.",
        "    - Par tienda-ítem faltante: ambos existen por separado, pero no juntos.",
        "",
        "VISTA HISTÓRICA (hojas 'Histórico - ...') — PROXY (aproximada)",
        "  Solo existe un Celes.xlsx (el de hoy); no hay históricos que reproducir.",
        "  Por eso el histórico se infiere desde los Outputs ya generados, con la",
        "  'firma sin datos': Consumo Diario = 0 Y Stock Antes = 0 Y Días Actual",
        "  vacío. Limitaciones:",
        "    - Subestima: los sin-match que NO recibieron pedido (el ítem no tenía",
        "      stock en CEDI ese día) no aparecen en la hoja Distribución. Ver la",
        "      hoja 'Histórico - Fechas': la columna Oficial (de la hoja Alertas de",
        "      cada Output) es mayor que el proxy por esta razón.",
        "    - Ruido: podría incluir alguna tienda que Celes SÍ reporta con rotación",
        "      e inventario genuinamente en 0 (indistinguible en el Output).",
        "  Los Outputs anteriores al 2026-07-02 no traen instrumentación de",
        "  sin-match, por lo que ahí no hay conteo Oficial con qué contrastar.",
        "",
        "RECOMENDACIÓN",
        "  Para tener el detalle exacto día a día sin auditar por fuera, se puede",
        "  añadir una hoja 'Sin Match Celes' al Output diario. Queda fuera de esta",
        "  entrega por decisión del negocio.",
    ]
    for n in notas:
        ws.append([n])
    ws.column_dimensions["A"].width = 82

    OUTPUT.mkdir(exist_ok=True)
    wb.save(AUDIT_PATH)


def main():
    print("=" * 70)
    print("AUDITORÍA — Combinaciones (tienda × ítem) SIN DATOS en Celes")
    print("=" * 70)

    print("\n[1/3] Vista HOY (reproduciendo el pipeline real)...")
    hoy_df, total_hoy = vista_hoy()
    n = len(hoy_df)
    print(f"  Combinaciones sin match en Celes HOY: {n} "
          f"({round(100 * n / total_hoy, 1) if total_hoy else 0}% de {total_hoy} pares elegibles)")
    if n:
        print(f"  Ítems distintos: {hoy_df['item_code'].nunique()} | "
              f"Tiendas distintas: {hoy_df['store_code'].nunique()}")
        print("  Causas:")
        for causa, c in hoy_df["causa"].value_counts().items():
            print(f"    - {causa}: {c}")
        print("  Top ítems hoy:")
        top = hoy_df.groupby(["item_code", "item_desc"]).size().sort_values(ascending=False).head(5)
        for (item, desc), c in top.items():
            print(f"    {int(item):>5}  {desc[:35]:<35} {c} tienda(s)")

    print("\n[2/3] Vista HISTÓRICA (proxy sobre Outputs)...")
    hist = vista_historica()
    n_pares = len(hist["sindatos"])
    n_occ = sum(len(v) for v in hist["sindatos"].values())
    n_fechas = len({f for v in hist["sindatos"].values() for f in v})
    print(f"  Outputs analizados: {len(hist['files'])}")
    print(f"  Pares (tienda×ítem) distintos con sin-datos: {n_pares}")
    print(f"  Ocurrencias totales (par×fecha): {n_occ} en {n_fechas} fechas")

    print("\n[3/3] Exportando Excel de auditoría...")
    exportar(hoy_df, total_hoy, hist)
    print(f"  Archivo generado: {AUDIT_PATH}")
    print("\nListo.")


if __name__ == "__main__":
    main()
