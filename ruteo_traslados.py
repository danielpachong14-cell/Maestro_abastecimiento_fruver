"""
ruteo_traslados.py
Módulo de ruteo geográfico para traslados de inventario.

Lee la hoja "Traslados Misma Zona" del Plan_Traslados más reciente y genera
un rutero detallado por zona usando OR-Tools (TSP con restricciones de
precedencia) con distancias OSRM (viales reales) o Haversine como fallback.

Ejecutar después de analisis_traslados.py:
    python ruteo_traslados.py
"""

from pathlib import Path
from datetime import date
import math
import json
import urllib.request
import pandas as pd

try:
    from ortools.constraint_solver import routing_enums_pb2, pywrapcp
    _ORTOOLS = True
except ImportError:
    _ORTOOLS = False
    print("[AVISO] OR-Tools no instalado. Ejecutar: pip install ortools")

# ─── Constantes configurables ────────────────────────────────────────────────
CEDI_LAT        = 4.8040    # Coordenadas del CEDI (depot neutro)
CEDI_LON        = -74.1042
INCLUIR_CEDI    = False     # True → la ruta empieza y termina en el CEDI
FACTOR_URBANO   = 1.30      # Penalización para Haversine (solo si USE_OSRM=False)
TIEMPO_LIMITE_S = 30        # Tiempo máximo del solver OR-Tools (segundos)
MAX_VUELTAS     = 10        # Máximo de pasadas por zona (válvula de seguridad)
USE_OSRM        = True      # True → distancias viales reales via OSRM; False → Haversine
INPUT_DIR       = Path("Input")
OUTPUT_DIR      = Path("Output")
# ─────────────────────────────────────────────────────────────────────────────


# ── Haversine (fallback) ──────────────────────────────────────────────────────

def haversine(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    """Gran distancia circular en metros."""
    R = 6_371_000
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    a = (
        math.sin(math.radians(lat2 - lat1) / 2) ** 2
        + math.cos(phi1) * math.cos(phi2)
        * math.sin(math.radians(lon2 - lon1) / 2) ** 2
    )
    return 2 * R * math.asin(math.sqrt(a))


# ── Matrices de distancias ────────────────────────────────────────────────────

def build_distance_matrix(nodes: list) -> list:
    """Matriz n×n en centímetros (Haversine × FACTOR_URBANO)."""
    n = len(nodes)
    matrix = [[0] * n for _ in range(n)]
    for i in range(n):
        if nodes[i]["role"] == "sink":
            continue
        for j in range(n):
            if i == j:
                continue
            if nodes[j]["role"] == "sink":
                matrix[i][j] = 0
            else:
                d = haversine(
                    nodes[i]["lat"], nodes[i]["lon"],
                    nodes[j]["lat"], nodes[j]["lon"],
                )
                matrix[i][j] = int(d * FACTOR_URBANO * 100)
    return matrix


def build_distance_matrix_osrm(nodes: list) -> list:
    """
    Matriz n×n en centímetros usando la API pública de OSRM
    (distancias viales reales en Colombia).
    Si OSRM no responde o falla, cae a Haversine × FACTOR_URBANO.
    """
    n = len(nodes)
    matrix = [[0] * n for _ in range(n)]

    # Solo nodos con coordenadas reales (excluir sink)
    real_idx = [i for i, nd in enumerate(nodes) if nd["role"] != "sink"]
    if len(real_idx) < 2:
        return build_distance_matrix(nodes)

    # OSRM usa formato lon,lat (¡no lat,lon!)
    coords_str = ";".join(
        f"{nodes[i]['lon']},{nodes[i]['lat']}"
        for i in real_idx
    )
    url = (
        f"http://router.project-osrm.org/table/v1/driving/{coords_str}"
        f"?annotations=distance"
    )

    try:
        req = urllib.request.Request(url, headers={"User-Agent": "ruteo_traslados/1.0"})
        with urllib.request.urlopen(req, timeout=30) as resp:
            data = json.loads(resp.read().decode("utf-8"))

        if data.get("code") != "Ok":
            raise ValueError(f"OSRM code={data.get('code')}")

        osrm = data["distances"]  # metros (float o None si sin ruta)

        for ii, i in enumerate(real_idx):
            for jj, j in enumerate(real_idx):
                if i == j:
                    continue
                d = osrm[ii][jj]
                if d is None:
                    # Sin ruta vial: fallback Haversine para este par
                    d = haversine(
                        nodes[i]["lat"], nodes[i]["lon"],
                        nodes[j]["lat"], nodes[j]["lon"],
                    ) * FACTOR_URBANO
                matrix[i][j] = int(d * 100)  # metros → cm

        return matrix

    except Exception as exc:
        print(f"  [AVISO] OSRM no disponible ({exc}). Usando Haversine x{FACTOR_URBANO}.")
        return build_distance_matrix(nodes)


# ── Topological nearest-neighbor ──────────────────────────────────────────────

def _topo_nn(nodes: list, pairs_node: list, depot_idx: int) -> list:
    """Ruteo topológico nearest-neighbor con garantía de precedencia."""
    from collections import defaultdict

    successors: dict = defaultdict(set)
    in_degree:  dict = defaultdict(int)

    real_nodes = {
        i for i, nd in enumerate(nodes)
        if nd["role"] not in ("sink", "CEDI")
    }

    for p, d in pairs_node:
        successors[p].add(d)
        in_degree[d] += 1

    available = {n for n in real_nodes if in_degree[n] == 0}
    remaining = set(real_nodes)
    route = [depot_idx]

    while remaining:
        current    = route[-1]
        candidates = available & remaining
        if not candidates:
            candidates = remaining

        best = min(
            candidates,
            key=lambda idx: haversine(
                nodes[current]["lat"], nodes[current]["lon"],
                nodes[idx]["lat"],    nodes[idx]["lon"],
            ),
        )

        route.append(best)
        remaining.discard(best)
        available.discard(best)

        for succ in successors[best]:
            if succ in remaining:
                in_degree[succ] -= 1
                if in_degree[succ] == 0:
                    available.add(succ)

    return route


# ── OR-Tools TSP ──────────────────────────────────────────────────────────────

def solve_tsp_precedence(
    nodes: list,
    pairs_node: list,
    depot_idx: int,
    end_idx: int,
    dist_matrix: list,
) -> list:
    """TSP con restricciones de precedencia. Cae a topo-NN si OR-Tools falla."""
    if not _ORTOOLS:
        return _topo_nn(nodes, pairs_node, depot_idx)

    n = len(nodes)
    manager = pywrapcp.RoutingIndexManager(n, 1, [depot_idx], [end_idx])
    routing  = pywrapcp.RoutingModel(manager)

    def dist_cb(from_idx, to_idx):
        return dist_matrix[manager.IndexToNode(from_idx)][manager.IndexToNode(to_idx)]

    transit_cb = routing.RegisterTransitCallback(dist_cb)
    routing.SetArcCostEvaluatorOfAllVehicles(transit_cb)

    def _unit_cb(from_idx, to_idx):
        return 1

    unit_cb = routing.RegisterTransitCallback(_unit_cb)
    routing.AddDimension(unit_cb, 0, n, True, "Position")
    pos_dim = routing.GetDimensionOrDie("Position")

    for p_node, d_node in pairs_node:
        p_idx = manager.NodeToIndex(p_node)
        d_idx = manager.NodeToIndex(d_node)
        routing.solver().Add(
            pos_dim.CumulVar(p_idx) <= pos_dim.CumulVar(d_idx)
        )

    params = pywrapcp.DefaultRoutingSearchParameters()
    params.first_solution_strategy = (
        routing_enums_pb2.FirstSolutionStrategy.GLOBAL_CHEAPEST_ARC
    )
    params.local_search_metaheuristic = (
        routing_enums_pb2.LocalSearchMetaheuristic.GUIDED_LOCAL_SEARCH
    )
    params.time_limit.seconds = TIEMPO_LIMITE_S

    solution = routing.SolveWithParameters(params)
    if not solution:
        return _topo_nn(nodes, pairs_node, depot_idx)

    route = []
    index = routing.Start(0)
    while not routing.IsEnd(index):
        route.append(manager.IndexToNode(index))
        index = solution.Value(routing.NextVar(index))
    route.append(end_idx)
    return route


# ── Filtro de pares acíclicos ─────────────────────────────────────────────────

def _build_acyclic_pairs(pairs_raw: list, bodega_to_node: dict) -> list:
    """Filtra pares cíclicos via DFS incremental. Devuelve solo pares acíclicos."""
    from collections import defaultdict

    graph: dict = defaultdict(set)

    def has_path(src: str, dst: str) -> bool:
        visited: set = set()
        stack = [src]
        while stack:
            node = stack.pop()
            if node == dst:
                return True
            if node in visited:
                continue
            visited.add(node)
            stack.extend(graph[node])
        return False

    acyclic = []
    for o, d in pairs_raw:
        if o == d:
            continue
        if not has_path(d, o):
            acyclic.append((bodega_to_node[o], bodega_to_node[d]))
            graph[o].add(d)

    return acyclic


# ── Una vuelta de ruteo ───────────────────────────────────────────────────────

def _route_single_pass(
    df_transfers: pd.DataFrame,
    store_info: dict,
    zona_label: str,
    _is_first_pass: bool,
) -> tuple:
    """
    Ejecuta una vuelta de ruteo para df_transfers.
    Retorna (rows, df_uncovered, n_stops, dist_total_km).
    df_uncovered: traslados donde origen quedo DESPUES del destino → otra vuelta.
    """
    unique_bodegas = sorted(
        set(df_transfers["Bodega Origen"].astype(str).str.strip()) |
        set(df_transfers["Bodega Destino"].astype(str).str.strip())
    )

    pairs_raw = list({
        (str(r["Bodega Origen"]).strip(), str(r["Bodega Destino"]).strip())
        for _, r in df_transfers[["Bodega Origen", "Bodega Destino"]].iterrows()
        if str(r["Bodega Origen"]).strip() != str(r["Bodega Destino"]).strip()
    })

    nodes: list = [{"lat": CEDI_LAT, "lon": CEDI_LON, "role": "CEDI", "bodega": "CEDI"}]
    b2n: dict = {}
    for b in unique_bodegas:
        b2n[b] = len(nodes)
        nodes.append({**store_info[b], "role": "store", "bodega": b})
    sink_idx = len(nodes)
    nodes.append({"lat": 0.0, "lon": 0.0, "role": "sink", "bodega": "_SINK"})

    pairs_node = _build_acyclic_pairs(pairs_raw, b2n)
    n_cyclic = len(pairs_raw) - len(pairs_node)

    print(f"\n  {zona_label}: {len(unique_bodegas)} tiendas | "
          f"{len(pairs_node)} restricciones ({n_cyclic} ciclicas)")

    if USE_OSRM:
        dist_matrix = build_distance_matrix_osrm(nodes)
    else:
        dist_matrix = build_distance_matrix(nodes)

    print(f"  Ejecutando OR-Tools para '{zona_label}'...")
    route = solve_tsp_precedence(nodes, pairs_node, 0, sink_idx, dist_matrix)

    route_clean = [
        idx for idx in route
        if nodes[idx]["role"] != "sink"
        and not (nodes[idx]["role"] == "CEDI" and not INCLUIR_CEDI)
    ]
    n_stops = len(route_clean)

    visit_pos: dict = {
        nodes[idx]["bodega"]: pos
        for pos, idx in enumerate(route_clean)
        if nodes[idx]["role"] == "store"
    }

    rows: list = []
    primera_fila = True
    dist_total = 0.0

    for list_pos, node_idx in enumerate(route_clean):
        nd     = nodes[node_idx]
        bodega = nd["bodega"]

        if list_pos < n_stops - 1:
            nxt_idx = route_clean[list_pos + 1]
            dist_km = round(dist_matrix[node_idx][nxt_idx] / 100_000, 2)
            dist_total += dist_km
        else:
            dist_km = ""

        parada_num = list_pos + 1

        if bodega == "CEDI":
            is_inicio = parada_num == 1
            rows.append({
                "Zona":              zona_label if primera_fila else "",
                "Parada":            parada_num,
                "Bodega":            "CEDI",
                "COD SIESA":         "—",
                "Nombre Tienda":     "Centro de Distribucion",
                "Accion":            "Inicio CEDI" if is_inicio else "Fin CEDI",
                "Cuadrante":         "",
                "Item":              "",
                "Producto":          "",
                "Cajas":             "",
                "Bodega Origen":     "",
                "Nombre Origen":     "",
                "Bodega Destino":    "",
                "Nombre Destino":    "",
                "Total Recoge":      "",
                "Total Entrega":     "",
                "Dist al Sig. (km)": dist_km,
            })
            primera_fila = False
            continue

        info   = store_info[bodega]
        cod    = info["cod_siesa"]
        nombre = info["nombre"]

        all_deliveries = df_transfers[df_transfers["Bodega Destino"] == bodega]
        deliveries = all_deliveries[
            all_deliveries["Bodega Origen"].map(
                lambda x: visit_pos.get(str(x).strip(), 9999) < list_pos
            )
        ]

        all_pickups = df_transfers[df_transfers["Bodega Origen"] == bodega]
        pickups = all_pickups[
            all_pickups["Bodega Destino"].map(
                lambda x: visit_pos.get(str(x).strip(), -1) > list_pos
            )
        ]

        total_recoge  = int(pickups["Cajas a Trasladar"].sum())    if not pickups.empty    else 0
        total_entrega = int(deliveries["Cajas a Trasladar"].sum()) if not deliveries.empty else 0

        ops: list = []
        for _, r in deliveries.iterrows():
            orig_b = str(r["Bodega Origen"]).strip()
            ops.append({
                "accion":         "ENTREGA",
                "cuadrante":      r["Cuadrante"],
                "item":           r["Item"],
                "producto":       r["Producto"],
                "cajas":          int(r["Cajas a Trasladar"]),
                "bodega_origen":  orig_b,
                "nombre_origen":  store_info[orig_b]["nombre"] if orig_b in store_info else orig_b,
                "bodega_destino": bodega,
                "nombre_destino": nombre,
            })
        for _, r in pickups.iterrows():
            dest_b = str(r["Bodega Destino"]).strip()
            ops.append({
                "accion":         "RECOGE",
                "cuadrante":      r["Cuadrante"],
                "item":           r["Item"],
                "producto":       r["Producto"],
                "cajas":          int(r["Cajas a Trasladar"]),
                "bodega_origen":  bodega,
                "nombre_origen":  nombre,
                "bodega_destino": dest_b,
                "nombre_destino": store_info[dest_b]["nombre"] if dest_b in store_info else dest_b,
            })

        if not ops:
            continue

        for op_i, op in enumerate(ops):
            primer_op = op_i == 0
            rows.append({
                "Zona":              zona_label if primera_fila else "",
                "Parada":            parada_num,
                "Bodega":            bodega,
                "COD SIESA":         cod,
                "Nombre Tienda":     nombre,
                "Accion":            op["accion"],
                "Cuadrante":         op["cuadrante"],
                "Item":              op["item"],
                "Producto":          op["producto"],
                "Cajas":             op["cajas"],
                "Bodega Origen":     op["bodega_origen"],
                "Nombre Origen":     op["nombre_origen"],
                "Bodega Destino":    op["bodega_destino"],
                "Nombre Destino":    op["nombre_destino"],
                "Total Recoge":      total_recoge  if primer_op else "",
                "Total Entrega":     total_entrega if primer_op else "",
                "Dist al Sig. (km)": dist_km       if primer_op else "",
            })
            primera_fila = False

    df_uncovered = df_transfers[
        df_transfers.apply(
            lambda r: visit_pos.get(str(r["Bodega Origen"]).strip(), -1) >
                      visit_pos.get(str(r["Bodega Destino"]).strip(), 9999),
            axis=1
        )
    ].copy()

    suffix = f" | {len(df_uncovered)} para sig. vuelta" if not df_uncovered.empty else ""
    print(f"  >> {n_stops} paradas | ~{dist_total:.1f} km{suffix}")

    return rows, df_uncovered, n_stops, dist_total


# ── Validación de cobertura ───────────────────────────────────────────────────

def _validate_coverage(
    df_plan: pd.DataFrame,
    all_rows: list,
    df_pendientes: pd.DataFrame,
) -> tuple:
    """
    Reconcilia cajas entre el plan de traslados y el rutero generado.
    Retorna (df_missing, summary_dict).
    """
    total_cajas_plan = int(df_plan["Cajas a Trasladar"].sum())

    recoge_keys: set  = set()
    entrega_keys: set = set()
    total_cajas_recoge  = 0
    total_cajas_entrega = 0

    for r in all_rows:
        if not isinstance(r.get("Cajas"), int):
            continue
        key = (
            str(r.get("Bodega Origen",  "")).strip(),
            str(r.get("Bodega Destino", "")).strip(),
            str(r.get("Item",           "")).strip(),
        )
        if r["Accion"] == "RECOGE":
            recoge_keys.add(key)
            total_cajas_recoge += r["Cajas"]
        elif r["Accion"] == "ENTREGA":
            entrega_keys.add(key)
            total_cajas_entrega += r["Cajas"]

    pendiente_keys: set = set()
    if not df_pendientes.empty:
        for _, r in df_pendientes.iterrows():
            pendiente_keys.add((
                str(r["Bodega Origen"]).strip(),
                str(r["Bodega Destino"]).strip(),
                str(r["Item"]).strip(),
            ))

    missing_rows = []
    for _, row in df_plan.iterrows():
        key = (
            str(row["Bodega Origen"]).strip(),
            str(row["Bodega Destino"]).strip(),
            str(row["Item"]).strip(),
        )
        has_recoge  = key in recoge_keys
        has_entrega = key in entrega_keys

        if has_recoge and has_entrega:
            continue

        if key in pendiente_keys:
            motivo = f"Pendiente (>{MAX_VUELTAS} vueltas)"
        elif not has_recoge and not has_entrega:
            motivo = "Sin RECOGE ni ENTREGA en rutero"
        elif not has_recoge:
            motivo = "Sin RECOGE en rutero"
        else:
            motivo = "Sin ENTREGA en rutero"

        missing_rows.append({
            "Zona":           row["Zona Origen"],
            "Bodega Origen":  row["Bodega Origen"],
            "Bodega Destino": row["Bodega Destino"],
            "Item":           row["Item"],
            "Producto":       row["Producto"],
            "Cajas":          int(row["Cajas a Trasladar"]),
            "Motivo":         motivo,
        })

    df_missing = pd.DataFrame(missing_rows)

    summary = {
        "total_plan":    total_cajas_plan,
        "total_recoge":  total_cajas_recoge,
        "total_entrega": total_cajas_entrega,
        "ok": (
            total_cajas_plan == total_cajas_recoge == total_cajas_entrega
            and df_missing.empty
        ),
    }

    return df_missing, summary


# ── Mapa HTML interactivo ─────────────────────────────────────────────────────

def _generar_mapa(all_rows: list, store_info: dict, out_path: Path) -> None:
    """Genera mapa HTML interactivo con rutas por zona y marcadores numerados."""
    try:
        import folium
    except ImportError:
        print("  [AVISO] folium no instalado: pip install folium")
        return

    COLORS = [
        "#e6194b", "#3cb44b", "#4363d8", "#f58231", "#911eb4",
        "#42d4f4", "#9a6324", "#469990", "#f032e6", "#808000",
    ]

    # Reconstruir paradas por zona (zona solo aparece en primera fila de cada bloque)
    zone_stops: dict = {}
    current_zona: str = ""

    for row in all_rows:
        if row.get("Zona"):
            current_zona = row["Zona"]
            if current_zona not in zone_stops:
                zone_stops[current_zona] = {}

        if not current_zona:
            continue

        bodega = str(row.get("Bodega", "")).strip()
        if bodega in ("CEDI", "_SINK", ""):
            continue

        parada = row.get("Parada", 0)
        if parada in zone_stops[current_zona]:
            continue

        info = store_info.get(bodega, {})
        lat, lon = info.get("lat", 0.0), info.get("lon", 0.0)
        if not lat or not lon:
            continue

        zone_stops[current_zona][parada] = {
            "bodega": bodega,
            "cod":    info.get("cod_siesa", ""),
            "nombre": row.get("Nombre Tienda", bodega),
            "lat":    lat,
            "lon":    lon,
            "tr":     row.get("Total Recoge",  ""),
            "te":     row.get("Total Entrega", ""),
        }

    if not zone_stops:
        print("  [AVISO] Sin datos para generar mapa.")
        return

    # Centro del mapa = centroide de todas las tiendas
    all_lats = [v["lat"] for d in zone_stops.values() for v in d.values()]
    all_lons = [v["lon"] for d in zone_stops.values() for v in d.values()]
    center   = [sum(all_lats) / len(all_lats), sum(all_lons) / len(all_lons)]

    m = folium.Map(location=center, zoom_start=12, tiles="CartoDB positron")

    # Color por zona base (Vuelta 1 y 2 comparten color, difieren en línea)
    base_zones = sorted({z.split(" - Vuelta")[0] for z in zone_stops})
    color_map  = {z: COLORS[i % len(COLORS)] for i, z in enumerate(base_zones)}

    for zona_label, stops_dict in sorted(zone_stops.items()):
        base_zone  = zona_label.split(" - Vuelta")[0]
        color      = color_map.get(base_zone, "#666666")
        is_vuelta2 = "Vuelta" in zona_label

        ordered = sorted(stops_dict.items())  # (parada_num, info)
        if not ordered:
            continue

        fg = folium.FeatureGroup(name=zona_label, show=True)

        # Línea de ruta: sólida para Vuelta 1, punteada para Vuelta 2+
        coords = [[v["lat"], v["lon"]] for _, v in ordered]
        folium.PolyLine(
            coords,
            color      = color,
            weight     = 2 if is_vuelta2 else 4,
            opacity    = 0.75,
            dash_array = "6 6" if is_vuelta2 else None,
            tooltip    = zona_label,
        ).add_to(fg)

        # Marcadores numerados con popup
        for parada, stop in ordered:
            gmaps = (
                f"https://www.google.com/maps/search/?api=1"
                f"&query={stop['lat']},{stop['lon']}"
            )
            popup_html = (
                f"<div style='font-family:sans-serif;font-size:13px;min-width:180px'>"
                f"<b>{stop['nombre']}</b><br>"
                f"<span style='color:#555'>{stop['bodega']} | {stop['cod']}</span><br>"
                f"<hr style='margin:4px 0'>"
                f"{zona_label} — Parada <b>{parada}</b><br>"
                f"RECOGE: <b style='color:#27ae60'>{stop['tr']}</b> cajas&nbsp;&nbsp;"
                f"ENTREGA: <b style='color:#e74c3c'>{stop['te']}</b> cajas<br>"
                f"<a href='{gmaps}' target='_blank' "
                f"style='font-size:11px;color:#2980b9'>Ver en Google Maps</a>"
                f"</div>"
            )
            folium.Marker(
                location = [stop["lat"], stop["lon"]],
                popup    = folium.Popup(popup_html, max_width=240),
                icon     = folium.DivIcon(
                    html = (
                        f'<div style="font-size:11px;font-weight:bold;color:white;'
                        f'background:{color};border-radius:50%;'
                        f'width:24px;height:24px;text-align:center;line-height:24px;'
                        f'border:2px solid white;'
                        f'box-shadow:0 2px 4px rgba(0,0,0,.5);">'
                        f'{parada}</div>'
                    ),
                    icon_size   = (24, 24),
                    icon_anchor = (12, 12),
                ),
            ).add_to(fg)

        fg.add_to(m)

    folium.LayerControl(collapsed=False).add_to(m)
    m.save(str(out_path))
    print(f"  Mapa  : {out_path.name}")


# ── Main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    # ── 1. Cargar DB_Tiendas ──────────────────────────────────────────────────
    tiendas_path = INPUT_DIR / "DB_Tiendas.xlsx"
    print(f"Cargando {tiendas_path.name}...")
    df_t = pd.read_excel(tiendas_path)
    df_t.columns = df_t.columns.str.strip()

    cod_col = "COD SIESA" if "COD SIESA" in df_t.columns else "C.O"
    for col in ("BODEGA", "NOMBRE DE TIENDA", "LATITUD", "LONGITUD", cod_col):
        if col not in df_t.columns:
            raise ValueError(
                f"Columna requerida '{col}' no encontrada en {tiendas_path.name}. "
                f"Disponibles: {list(df_t.columns)}"
            )

    store_info: dict = {}
    n_corregidas = 0
    for _, row in df_t.iterrows():
        if pd.isna(row["BODEGA"]):
            continue
        bodega = str(row["BODEGA"]).strip()

        lat, lon = None, None

        # 1. Fuente primaria: columna COORDENADAS (copiada de Google Maps)
        coord_raw = str(row.get("COORDENADAS", "")).strip()
        clat, clon = None, None
        if coord_raw and coord_raw != "nan":
            try:
                parts = coord_raw.split(",")
                clat = float(parts[0].strip())
                clon = float(parts[1].strip())
                if 0 <= clat <= 15 and -82 <= clon <= -65:
                    lat, lon = clat, clon
                else:
                    clat, clon = None, None
            except (ValueError, IndexError):
                pass

        # 2. Fallback: LATITUD / LONGITUD con correcciones de signo/swap
        llat, llon = None, None
        try:
            llat = float(str(row["LATITUD"]).split(",")[0].strip())
            llon = float(str(row["LONGITUD"]).split(",")[0].strip())
            if llat < -10 and 0 < llon < 15:
                llat, llon = llon, llat
                n_corregidas += 1
            if llon > 60:
                llon = -llon
                n_corregidas += 1
            if not (0 <= llat <= 15 and -82 <= llon <= -65):
                llat, llon = None, None
        except (ValueError, TypeError):
            llat, llon = None, None

        if lat is None:
            if llat is not None:
                lat, lon = llat, llon
            else:
                continue
        elif llat is not None and (abs(clat - llat) > 0.5 or abs(clon - llon) > 0.5):
            # Ambas fuentes válidas pero difieren >0.5° (~55 km) — posible error en COORDENADAS
            nombre = str(row.get("NOMBRE DE TIENDA", "")).strip()
            print(f"  [COORD] {bodega} ({nombre}): "
                  f"COORDENADAS=({clat:.4f},{clon:.4f}) vs "
                  f"LAT/LON=({llat:.4f},{llon:.4f}) — verificar en DB_Tiendas.xlsx")

        store_info[bodega] = {
            "cod_siesa": row[cod_col],
            "nombre":    row["NOMBRE DE TIENDA"],
            "lat":       lat,
            "lon":       lon,
        }

    print(f"  {len(store_info)} tiendas con coordenadas"
          + (f" ({n_corregidas} correcciones lat/lon aplicadas)." if n_corregidas else "."))

    # ── 2. Plan de traslados más reciente ─────────────────────────────────────
    plan_files = sorted(OUTPUT_DIR.glob("Plan_Traslados_*.xlsx"), reverse=True)
    if not plan_files:
        raise FileNotFoundError(
            f"No hay archivos Plan_Traslados_*.xlsx en {OUTPUT_DIR}. "
            "Ejecutar primero analisis_traslados.py."
        )

    plan_path = plan_files[0]
    print(f"Cargando {plan_path.name} (hoja 'Traslados Misma Zona')...")
    try:
        df_tr = pd.read_excel(plan_path, sheet_name="Traslados Misma Zona")
    except Exception as exc:
        raise ValueError(
            f"No se pudo leer 'Traslados Misma Zona' de {plan_path.name}: {exc}"
        ) from exc

    if df_tr.empty:
        print("La hoja 'Traslados Misma Zona' esta vacia — sin traslados para rutear.")
        return

    df_tr.columns = df_tr.columns.str.strip()

    COLS_NEEDED = [
        "Bodega Origen", "Bodega Destino",
        "Item", "Producto", "Cuadrante", "Cajas a Trasladar", "Zona Origen",
    ]
    missing_cols = [c for c in COLS_NEEDED if c not in df_tr.columns]
    if missing_cols:
        raise ValueError(
            f"Columnas faltantes en 'Traslados Misma Zona': {missing_cols}. "
            f"Disponibles: {list(df_tr.columns)}"
        )

    df_tr = df_tr[COLS_NEEDED].dropna(subset=["Bodega Origen", "Bodega Destino"]).copy()
    df_tr["Bodega Origen"]  = df_tr["Bodega Origen"].astype(str).str.strip()
    df_tr["Bodega Destino"] = df_tr["Bodega Destino"].astype(str).str.strip()
    df_tr["Zona Origen"]    = df_tr["Zona Origen"].astype(str).str.strip()

    mask_ok = df_tr["Bodega Origen"].isin(store_info) & df_tr["Bodega Destino"].isin(store_info)
    skipped_bodegas: set = set()
    for _, row in df_tr[~mask_ok].iterrows():
        for b in [row["Bodega Origen"], row["Bodega Destino"]]:
            if b not in store_info:
                skipped_bodegas.add(b)
    df_tr = df_tr[mask_ok].copy()

    if skipped_bodegas:
        print(f"  [AVISO] {len(skipped_bodegas)} tiendas sin coordenadas omitidas: "
              f"{sorted(skipped_bodegas)}")
    if df_tr.empty:
        print("No hay traslados validos con coordenadas. Abortando.")
        return

    n_zonas = df_tr["Zona Origen"].nunique()
    print(f"  {len(df_tr)} traslados en {n_zonas} zona(s).")

    # ── 3. Ruteo N-pass por zona ──────────────────────────────────────────────
    all_rows: list        = []
    zonas                 = sorted(df_tr["Zona Origen"].dropna().unique())
    totales_zona: list    = []
    df_pendientes_list: list = []

    for zona in zonas:
        df_zona      = df_tr[df_tr["Zona Origen"] == zona].copy()
        df_remaining = df_zona.copy()
        vuelta       = 1

        while not df_remaining.empty and vuelta <= MAX_VUELTAS:
            label      = zona if vuelta == 1 else f"{zona} - Vuelta {vuelta}"
            prev_count = len(df_remaining)

            new_rows, df_remaining, n_stops, dist_km = _route_single_pass(
                df_remaining, store_info, label, _is_first_pass=(vuelta == 1)
            )
            all_rows.extend(new_rows)

            if n_stops > 0:
                totales_zona.append((label, n_stops, dist_km))

            if len(df_remaining) >= prev_count:
                break

            vuelta += 1

        if not df_remaining.empty:
            df_pendientes_list.append(df_remaining)
            print(f"  [AVISO] {len(df_remaining)} traslados de '{zona}' "
                  f"sin cubrir tras {MAX_VUELTAS} vueltas.")

    df_pendientes = (
        pd.concat(df_pendientes_list, ignore_index=True)
        if df_pendientes_list else pd.DataFrame()
    )

    # ── 4. Validación de cobertura ────────────────────────────────────────────
    df_missing, summary = _validate_coverage(df_tr, all_rows, df_pendientes)

    # ── 5. Guardar Excel ──────────────────────────────────────────────────────
    if not all_rows:
        print("No se generaron filas. Abortando.")
        return

    from openpyxl.styles import PatternFill, Font
    from openpyxl.utils import get_column_letter

    df_ruta = pd.DataFrame(all_rows)
    OUTPUT_DIR.mkdir(exist_ok=True)
    fecha_str = date.today().strftime("%Y%m%d")
    out_path  = OUTPUT_DIR / f"Rutero_{fecha_str}.xlsx"

    COL_WIDTHS = {
        "Zona":              20,
        "Parada":             8,
        "Bodega":            10,
        "COD SIESA":         10,
        "Nombre Tienda":     30,
        "Accion":            10,
        "Cuadrante":         20,
        "Item":              12,
        "Producto":          35,
        "Cajas":              8,
        "Bodega Origen":     14,
        "Nombre Origen":     30,
        "Bodega Destino":    14,
        "Nombre Destino":    30,
        "Total Recoge":      13,
        "Total Entrega":     13,
        "Dist al Sig. (km)": 16,
    }

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # ── Hoja "Ruta Optima" ────────────────────────────────────────────────
        df_ruta.to_excel(writer, sheet_name="Ruta Optima", index=False)
        ws = writer.sheets["Ruta Optima"]

        for col_i, col_name in enumerate(df_ruta.columns, start=1):
            ws.column_dimensions[get_column_letter(col_i)].width = COL_WIDTHS.get(col_name, 15)

        fill_v1 = PatternFill("solid", fgColor="D9D9D9")  # gris  — vuelta 1
        fill_vn = PatternFill("solid", fgColor="BDD7EE")  # azul  — vueltas 2+
        bold    = Font(bold=True)
        n_cols  = len(df_ruta.columns)
        col_idx = {name: i + 1 for i, name in enumerate(df_ruta.columns)}

        for row_i, row_data in enumerate(all_rows, start=2):
            zona_val = row_data["Zona"]
            if zona_val:
                fill = fill_vn if "Vuelta" in str(zona_val) else fill_v1
                for c in range(1, n_cols + 1):
                    cell = ws.cell(row=row_i, column=c)
                    cell.fill = fill
                    cell.font = bold

            for field in ("Total Recoge", "Total Entrega"):
                if row_data[field] != "":
                    ws.cell(row=row_i, column=col_idx[field]).font = bold

        # ── Hoja "Validacion" ─────────────────────────────────────────────────
        ws_val   = writer.book.create_sheet("Validacion")
        fill_ok  = PatternFill("solid", fgColor="C6EFCE")
        fill_err = PatternFill("solid", fgColor="FFC7CE")

        for col_i, header in enumerate(("Concepto", "Cajas", "Estado"), start=1):
            ws_val.cell(row=1, column=col_i, value=header).font = bold

        rows_summary = [
            ("Total cajas en Plan_Traslados", summary["total_plan"]),
            ("Total cajas RECOGE en Rutero",  summary["total_recoge"]),
            ("Total cajas ENTREGA en Rutero", summary["total_entrega"]),
        ]
        for row_i, (concepto, cajas) in enumerate(rows_summary, start=2):
            ws_val.cell(row=row_i, column=1, value=concepto)
            ws_val.cell(row=row_i, column=2, value=cajas)
            ok_cell = ws_val.cell(
                row=row_i, column=3,
                value="OK" if cajas == summary["total_plan"] else "DIFERENCIA"
            )
            ok_cell.fill = fill_ok if cajas == summary["total_plan"] else fill_err

        status_cell = ws_val.cell(
            row=5, column=1,
            value="COBERTURA COMPLETA" if summary["ok"] else "HAY TRASLADOS SIN CUBRIR"
        )
        status_cell.font = Font(bold=True, size=14)
        status_cell.fill = fill_ok if summary["ok"] else fill_err

        ws_val.column_dimensions["A"].width = 38
        ws_val.column_dimensions["B"].width = 14
        ws_val.column_dimensions["C"].width = 14

        if not df_missing.empty:
            ws_val.cell(row=7, column=1, value="Traslados sin cobertura completa:").font = bold
            headers = list(df_missing.columns)
            for col_i, h in enumerate(headers, start=1):
                ws_val.cell(row=8, column=col_i, value=h).font = bold
            for row_i, (_, mrow) in enumerate(df_missing.iterrows(), start=9):
                for col_i, h in enumerate(headers, start=1):
                    ws_val.cell(row=row_i, column=col_i, value=mrow[h])
            for col_i, h in enumerate(headers, start=1):
                ws_val.column_dimensions[get_column_letter(col_i)].width = max(len(h) + 2, 14)

    # ── 6. Generar mapa HTML ──────────────────────────────────────────────────
    mapa_path = OUTPUT_DIR / f"Mapa_Rutero_{fecha_str}.html"
    _generar_mapa(all_rows, store_info, mapa_path)

    # ── 7. Resumen en consola ─────────────────────────────────────────────────
    dist_total  = sum(d for _, _, d in totales_zona)
    stops_total = sum(n for _, n, _ in totales_zona)
    dist_label  = "OSRM (viales reales)" if USE_OSRM else f"Haversine x{FACTOR_URBANO}"

    print(f"\n{'='*57}")
    print(f"  RUTERO GENERADO EXITOSAMENTE")
    print(f"{'='*57}")
    print(f"  Excel : {out_path.name}")
    print(f"  Mapa  : {mapa_path.name}")
    print(f"  Carpeta: {out_path.parent.resolve()}")
    print(f"  {'Zona':<28} {'Paradas':>8}  {'Dist (km)':>10}")
    print(f"  {'-'*50}")
    for label, n_stops, dist in totales_zona:
        print(f"  {label:<28} {n_stops:>8}  {dist:>9.1f}")
    print(f"  {'-'*50}")
    print(f"  {'TOTAL':<28} {stops_total:>8}  {dist_total:>9.1f}")
    print(f"  Distancias: {dist_label}")
    print(f"{'='*57}")

    print(f"\n  Validacion de cobertura:")
    print(f"  Plan:    {summary['total_plan']:>6} cajas")
    ok_r = summary["total_recoge"]  == summary["total_plan"]
    ok_e = summary["total_entrega"] == summary["total_plan"]
    print(f"  RECOGE:  {summary['total_recoge']:>6} cajas  {'[OK]' if ok_r else '[DIFERENCIA]'}")
    print(f"  ENTREGA: {summary['total_entrega']:>6} cajas  {'[OK]' if ok_e else '[DIFERENCIA]'}")
    if df_missing.empty:
        print(f"  Cobertura: COMPLETA")
    else:
        print(f"  Cobertura: {len(df_missing)} traslados sin cubrir (ver hoja 'Validacion')")
    print(f"{'='*57}\n")


if __name__ == "__main__":
    main()
