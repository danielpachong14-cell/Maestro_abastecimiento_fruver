"""Genera el Excel de salida formateado con la distribución.

Hojas:
  - Distribución:        el pedido por tienda-ítem (lo que se ejecuta).
  - Resumen:             totales de cajas por tienda.
  - Resumen Ítems:       totales y desglose por fase de distribución, agrupado por ítem.
  - Análisis Comprador:  por ítem, compara cajas disponibles en CEDI contra el mínimo
                         necesario para cubrir TARGET_DAYS en todas las tiendas, y
                         clasifica el nivel de riesgo para orientar la decisión de compra.
  - Alertas:              casos borde detectados durante el procesamiento (ítems sin
                         distribuir, combinaciones sin match en Celes, validación
                         cruzada con Portafolio, etc.). "OK — Sin alertas" si no hubo.
"""

import io

import numpy as np
import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from .algorithm import MIN_CAJAS_INICIAL, TARGET_DAYS, TOPE_EXCEDENTE

HEADER_FILL = PatternFill("solid", fgColor="2D6A4F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_ALIGN = Alignment(horizontal='center')

# Semáforo para '# / % Tiendas > Tope Excedente': mismos cortes que _nivel
# (>=50% alto, >=20% medio, >0 bajo, 0 ok) para que el color sea consistente
# con la clasificación de 'Nivel de Riesgo'.
_FILL_TOPE_ALTO  = PatternFill("solid", fgColor="F4A6A6")  # rojo claro
_FILL_TOPE_MEDIO = PatternFill("solid", fgColor="FBE0A6")  # naranja claro
_FILL_TOPE_BAJO  = PatternFill("solid", fgColor="FFF6B3")  # amarillo claro
_FILL_TOPE_OK    = PatternFill("solid", fgColor="D4EDDA")  # verde claro

_TOPE_DESCRIPCION = (
    '% de tiendas activas del ítem cuyo inventario proyectado tras el pedido '
    'supera el tope de excedente que el propio algoritmo usa para limitar el '
    'reparto del sobrante ({tope} días). A más alto el porcentaje, más cajas '
    'quedarían acumuladas en tienda por encima de ese límite.\n\n'
    'Verde = 0% · Amarillo = bajo (>0%) · Naranja = medio (≥20%) · Rojo = alto (≥50%).'
)

# Alertas ordenadas por severidad al escribir la hoja "Alertas".
_ORDEN_TIPO_ALERTA = {'CRÍTICO': 0, 'ERROR': 0, 'ADVERTENCIA': 1}


def _style_header(ws):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN


def _autofit(ws):
    for col in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in col), default=0) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max_len, 50)


def _resaltar_tope_excedente(ws, comprador_out):
    """Colorea '# / % Tiendas > Tope Excedente' por severidad y describe el
    análisis en un comentario sobre el encabezado de la columna de porcentaje."""
    col_n   = comprador_out.columns.get_loc(f'# Tiendas > Tope Excedente ({TOPE_EXCEDENTE}d)') + 1
    col_pct = comprador_out.columns.get_loc('% Tiendas > Tope Excedente') + 1

    for row_idx, pct in enumerate(comprador_out['% Tiendas > Tope Excedente'], start=2):
        if pct >= 50:
            fill = _FILL_TOPE_ALTO
        elif pct >= 20:
            fill = _FILL_TOPE_MEDIO
        elif pct > 0:
            fill = _FILL_TOPE_BAJO
        else:
            fill = _FILL_TOPE_OK
        ws.cell(row=row_idx, column=col_n).fill = fill
        ws.cell(row=row_idx, column=col_pct).fill = fill

    comentario = Comment(_TOPE_DESCRIPCION.format(tope=TOPE_EXCEDENTE), 'Distribuidor Fruver')
    comentario.width = 320
    comentario.height = 140
    ws.cell(row=1, column=col_pct).comment = comentario


def _build_alertas_sheet(alertas):
    """Arma el DataFrame de la hoja "Alertas". Los dicts de alerta traen
    claves heterogéneas según el tipo (`cajas_sin_distribuir`,
    `combos_sin_match`, `tiendas_sin_consumo`, ...) — `pd.DataFrame(alertas)`
    ya rellena con NaN las que falten por fila; solo se reordenan las
    columnas núcleo (Tipo/Ítem/Mensaje) primero y se ordenan por severidad."""
    if not alertas:
        return pd.DataFrame([{'Tipo': 'OK', 'Ítem': '-', 'Mensaje': 'Sin alertas'}])

    alertas_out = pd.DataFrame(alertas)
    alertas_out['_orden'] = alertas_out['tipo'].map(_ORDEN_TIPO_ALERTA).fillna(2)
    alertas_out = alertas_out.sort_values('_orden', kind='stable').drop(columns=['_orden'])

    core_cols = ['tipo', 'item', 'mensaje']
    extra_cols = [c for c in alertas_out.columns if c not in core_cols]
    alertas_out = alertas_out[core_cols + extra_cols].reset_index(drop=True)
    alertas_out.columns = ['Tipo', 'Ítem', 'Mensaje'] + [c.replace('_', ' ').title() for c in extra_cols]
    return alertas_out


def generate_excel(df_output, alertas=None, df_merged=None):
    """Genera el archivo Excel de distribución y lo retorna como bytes.

    df_merged (opcional): DataFrame pre-algoritmo con TODAS las tiendas elegibles.
    Si se provee, se genera la hoja 'Análisis Comprador'.
    """
    alertas = alertas or []
    has_data = df_output is not None and len(df_output) > 0

    # Pre-computar columnas derivadas una sola vez para reutilizar en todas las hojas.
    if has_data:
        d = df_output.copy()
        for col in ('cajas_fase1', 'cajas_fase2', 'cajas_fase3', 'vida_util'):
            if col not in d.columns:
                d[col] = 0
        d['_stock_despues'] = d['inventario_efectivo'] + d['cajas_asignadas']
        d['_dias_despues'] = np.where(
            d['consumo_diario'] > 0,
            (d['_stock_despues'] / d['consumo_diario']).round(1),
            np.nan,
        )
        # Umbral de riesgo de merma: min(10, vida_util) si vida_util > 0; si no, 10.
        # Se usa en "Análisis Comprador" para las estadísticas de riesgo por ítem.
        vida_util_num = pd.to_numeric(d['vida_util'], errors='coerce')
        d['_umbral'] = np.where(vida_util_num > 0, vida_util_num.clip(upper=10.0), 10.0)
    else:
        d = None

    # ── Hoja Distribución ──────────────────────────────────────────────────────
    if not has_data:
        dist_df = pd.DataFrame(columns=[
            'Centro Operacional de la Bodega', 'Nombre Tienda', 'Zona',
            'Código de Producto', 'Nombre Ítem', 'UM',
            'Consumo Diario', 'Stock Antes Pedido', 'Días Inventario Actual',
            'Pedido Final', 'Stock Después Pedido', 'Días Inventario Proyectado',
        ])
    else:
        dist_df = pd.DataFrame({
            'Centro Operacional de la Bodega': d['store_code'],
            'Nombre Tienda':                   d['store_name'],
            'Zona':                            d['zona'],
            'Código de Producto':              d['item_code'],
            'Nombre Ítem':                     d['item_desc'],
            'UM':                              d['um'],
            'Consumo Diario':                  d['consumo_diario'].round(2),
            'Stock Antes Pedido':              d['inventario_efectivo'].round(2),
            # dias_proyectados puede ser inf (consumo_diario==0 antes del pedido) —
            # se sanitiza a NaN (celda vacía) igual que la columna hermana de abajo,
            # en vez de escribir el string literal 'inf' en el Excel.
            'Días Inventario Actual':          d['dias_proyectados'].replace([np.inf, -np.inf], np.nan).round(1),
            'Pedido Final':                    d['cajas_asignadas'].astype(int),
            'Stock Después Pedido':            d['_stock_despues'].round(2),
            'Días Inventario Proyectado':      d['_dias_despues'],
        }).sort_values(['Centro Operacional de la Bodega', 'Código de Producto'])

    # ── Hoja Resumen ───────────────────────────────────────────────────────────
    if len(dist_df) > 0:
        resumen_df = (
            dist_df.groupby(
                ['Centro Operacional de la Bodega', 'Nombre Tienda', 'Zona'], as_index=False
            )['Pedido Final']
            .sum()
            .rename(columns={'Pedido Final': 'Total Cajas'})
            .sort_values('Total Cajas', ascending=False)
        )
    else:
        resumen_df = pd.DataFrame(
            columns=['Centro Operacional de la Bodega', 'Nombre Tienda', 'Zona', 'Total Cajas']
        )

    # ── Hoja Resumen Ítems ─────────────────────────────────────────────────────
    if has_data:
        resumitem_df = (
            d.groupby('item_code', as_index=False)
            .agg(
                item_desc=('item_desc', 'first'),
                total_cajas=('cajas_asignadas', 'sum'),
                fase1=('cajas_fase1', 'sum'),
                fase2=('cajas_fase2', 'sum'),
                fase3=('cajas_fase3', 'sum'),
            )
            .rename(columns={
                'item_code':   'Código Ítem',
                'item_desc':   'Nombre Ítem',
                'total_cajas': 'Total Cajas',
                'fase1':       'Cajas Fase 1',
                'fase2':       'Cajas Fase 2',
                'fase3':       'Cajas Fase 3',
            })
            .sort_values('Código Ítem')
        )
    else:
        resumitem_df = pd.DataFrame(columns=[
            'Código Ítem', 'Nombre Ítem', 'Total Cajas',
            'Cajas Fase 1', 'Cajas Fase 2', 'Cajas Fase 3',
        ])

    # ── Hoja Análisis Comprador ────────────────────────────────────────────────
    # Compara las cajas disponibles en CEDI contra el mínimo para cubrir TARGET_DAYS
    # en todas las tiendas, e identifica el exceso que el comprador puede reducir
    # sin afectar la cobertura mínima de inventario en tienda.
    _COMPRADOR_COLS = [
        'Código Ítem', 'Nombre Ítem', 'Cajas CEDI', '# Tiendas Activas',
        'Cajas Necesarias (3 días)', 'Exceso en CEDI (Cj)', '% Reducción Posible',
        '# Tiendas con Riesgo Merma', '% Tiendas con Riesgo',
        f'# Tiendas > Tope Excedente ({TOPE_EXCEDENTE}d)', '% Tiendas > Tope Excedente',
        'Exceso Cajas Est. Total', 'Nivel de Riesgo', 'Recomendación Compra',
    ]
    if has_data and df_merged is not None:
        mn = df_merged.copy()
        # Cajas mínimas que cada tienda necesita para alcanzar TARGET_DAYS de cobertura.
        necesidad_min = (TARGET_DAYS * mn['consumo_diario'] - mn['inventario_efectivo']).clip(lower=0)
        mn['_need'] = np.where(mn['consumo_diario'] > 0, np.ceil(necesidad_min), 0).astype(int)
        item_base = mn.groupby('item_code', as_index=False).agg(
            item_desc       =('item_desc',              'first'),
            cajas_cedi      =('cajas_disponibles_cedi', 'first'),
            n_tiendas       =('store_code',             'count'),
            cajas_min_3dias =('_need',                  'sum'),
        )
        item_base['exceso_cedi'] = (
            item_base['cajas_cedi'] - item_base['cajas_min_3dias']
        ).clip(lower=0)
        item_base['pct_reduccion'] = np.where(
            item_base['cajas_cedi'] > 0,
            (item_base['exceso_cedi'] / item_base['cajas_cedi'] * 100).round(1),
            0.0,
        )

        # Estadísticas de merma por ítem (d['_umbral'] ya fue calculado en la sección anterior).
        d['_is_merma'] = (
            (d['_dias_despues'].notna() & (d['_dias_despues'] > d['_umbral']))
            | (d['_dias_despues'].isna() & (d['cajas_asignadas'] > 0))
        ).astype(int)
        # Cajas en exceso por fila: para consumo=0 todas las cajas asignadas son exceso.
        d['_exceso_cajas_item'] = np.where(
            d['_dias_despues'].notna() & (d['consumo_diario'] > 0),
            ((d['_dias_despues'] - d['_umbral']) * d['consumo_diario']).clip(lower=0),
            d['cajas_asignadas'].astype(float),
        )
        # Mismo chequeo que _is_merma, pero contra TOPE_EXCEDENTE (días) en vez del
        # umbral por vida útil — identifica tiendas-ítem que quedan por encima del
        # tope que el propio algoritmo usa para limitar el reparto del sobrante.
        d['_excede_tope'] = (
            (d['_dias_despues'].notna() & (d['_dias_despues'] > TOPE_EXCEDENTE))
            | (d['_dias_despues'].isna() & (d['cajas_asignadas'] > 0))
        ).astype(int)

        item_merma = d.groupby('item_code', as_index=False).agg(
            n_tiendas_merma    =('_is_merma',          'sum'),
            exceso_cajas_total =('_exceso_cajas_item',  'sum'),
            n_tiendas_tope     =('_excede_tope',        'sum'),
        )

        cb = item_base.merge(item_merma, on='item_code', how='left')
        cb['n_tiendas_merma']    = cb['n_tiendas_merma'].fillna(0).astype(int)
        cb['exceso_cajas_total'] = cb['exceso_cajas_total'].fillna(0).round(1)
        cb['pct_tiendas_merma']  = (cb['n_tiendas_merma'] / cb['n_tiendas'] * 100).round(1)
        cb['n_tiendas_tope']     = cb['n_tiendas_tope'].fillna(0).astype(int)
        cb['pct_tiendas_tope']   = (cb['n_tiendas_tope'] / cb['n_tiendas'] * 100).round(1)

        def _nivel(row):
            if row['exceso_cedi'] <= 0:
                return 'OK'
            pct = row['pct_tiendas_merma']
            if pct >= 50:
                return 'ALTO'
            if pct >= 20:
                return 'MEDIO'
            if row['n_tiendas_merma'] > 0:
                return 'BAJO'
            return 'SIN RIESGO MERMA'

        cb['Nivel de Riesgo'] = cb.apply(_nivel, axis=1)

        def _recomendar(row):
            nivel = row['Nivel de Riesgo']
            exc   = int(row['Exceso en CEDI (Cj)'])
            min3  = int(row['Cajas Necesarias (3 días)'])
            if nivel in ('ALTO', 'MEDIO'):
                return f'Reducir {exc} Cj — mín. necesario: {min3} Cj'
            if nivel in ('BAJO', 'SIN RIESGO MERMA'):
                return f'Evaluar reducir {exc} Cj (bajo riesgo de merma)'
            return 'Compra adecuada o insuficiente'

        comprador_out = pd.DataFrame({
            'Código Ítem':                cb['item_code'],
            'Nombre Ítem':                cb['item_desc'],
            'Cajas CEDI':                 cb['cajas_cedi'],
            '# Tiendas Activas':          cb['n_tiendas'],
            'Cajas Necesarias (3 días)':  cb['cajas_min_3dias'],
            'Exceso en CEDI (Cj)':        cb['exceso_cedi'],
            '% Reducción Posible':        cb['pct_reduccion'],
            '# Tiendas con Riesgo Merma': cb['n_tiendas_merma'],
            '% Tiendas con Riesgo':       cb['pct_tiendas_merma'],
            f'# Tiendas > Tope Excedente ({TOPE_EXCEDENTE}d)': cb['n_tiendas_tope'],
            '% Tiendas > Tope Excedente': cb['pct_tiendas_tope'],
            'Exceso Cajas Est. Total':    cb['exceso_cajas_total'],
            'Nivel de Riesgo':            cb['Nivel de Riesgo'],
        })
        comprador_out['Recomendación Compra'] = comprador_out.apply(_recomendar, axis=1)

        _riesgo_order = {'ALTO': 0, 'MEDIO': 1, 'BAJO': 2, 'SIN RIESGO MERMA': 3, 'OK': 4}
        comprador_out['_sort_key'] = comprador_out['Nivel de Riesgo'].map(_riesgo_order)
        comprador_out = (
            comprador_out
            .sort_values(['_sort_key', 'Exceso en CEDI (Cj)'], ascending=[True, False])
            .drop(columns=['_sort_key'])
            .reset_index(drop=True)
        )
    else:
        comprador_out = None

    # ── Hoja Alertas ───────────────────────────────────────────────────────────
    alertas_out = _build_alertas_sheet(alertas)

    # ── Escritura ──────────────────────────────────────────────────────────────
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        dist_df.to_excel(writer, index=False, sheet_name='Distribución')
        resumen_df.to_excel(writer, index=False, sheet_name='Resumen')
        resumitem_df.to_excel(writer, index=False, sheet_name='Resumen Ítems', startrow=4)
        if comprador_out is not None:
            comprador_out.to_excel(writer, index=False, sheet_name='Análisis Comprador')
        alertas_out.to_excel(writer, index=False, sheet_name='Alertas')

        # Notas de fases en las primeras filas de "Resumen Ítems".
        _notas = [
            'Fase 1 — Urgentes: cubre tiendas AGOTADAS y en STOCK DE SEGURIDAD con cap '
            f'creciente por rondas (máx. {MIN_CAJAS_INICIAL} cajas/ronda). Todas reciben su 1.ª '
            'caja antes de que alguna reciba su 2.ª.',
            f'Fase 2 — Proporcional: completa hasta {TARGET_DAYS:g} días de inventario. Si el '
            'stock no alcanza para todos, la escasez se reparte de forma proporcional '
            '(corrección Hamilton para residuos enteros).',
            f'Fase 3 — Excedente: el sobrante se distribuye por rondas proporcional a '
            f'consumo_diario / días_actuales entre tiendas elegibles (consumo > 0, tope '
            f'{TOPE_EXCEDENTE} días). Las tiendas más atrasadas en días ganan más excedente, '
            'cerrando la brecha en vez de mantenerla.',
        ]
        ws_ri = writer.sheets['Resumen Ítems']
        for i, nota in enumerate(_notas, start=1):
            ws_ri.cell(row=i, column=1, value=nota)

        for sheet in ('Distribución', 'Resumen', 'Alertas'):
            ws = writer.sheets[sheet]
            _style_header(ws)
            _autofit(ws)

        if comprador_out is not None and 'Análisis Comprador' in writer.sheets:
            ws_cb = writer.sheets['Análisis Comprador']
            _style_header(ws_cb)
            _autofit(ws_cb)
            _resaltar_tope_excedente(ws_cb, comprador_out)

        # Encabezado de "Resumen Ítems" desplazado por las notas (fila 5 en Excel).
        for cell in ws_ri[5]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = HEADER_ALIGN
        _autofit(ws_ri)

    return buffer.getvalue()
